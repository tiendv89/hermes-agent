"""Tests for the document tools' storage-service integration.

Covers:
  - read_file: proxies to storage-service (mocked)
  - read_file: status.yaml → synthesized from workflow-backend feature detail
  - write_product_spec: proxies to storage-service
  - write_technical_design: proxies to storage-service
  - edit_document: read+apply+write via storage-service
  - storage_service_client unit tests (HTTP-level)
"""

from __future__ import annotations

import asyncio
import importlib
import importlib.util
import sys
import types
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

import pytest

REPO_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO_ROOT))


# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

_STORAGE_URL = "http://storage-service:8090"
_STORAGE_TOKEN = "storage-token-abc"
_WORKSPACE_ID = "ws-test"
_FEATURE_ID = "my-feature"
_OWNER = "testorg"
_REPO = "testws"
_VERSION_ID = "v1-uuid-abc"
_PRODUCT_SPEC_CONTENT = "# Product Spec\n\nContent here.\n"
_TD_CONTENT = "# Technical Design\n\nDesign here.\n"


def _make_workspace_context():
    return {
        "management_repo": _REPO,
        "repos": [{"id": _REPO, "github": f"https://github.com/{_OWNER}/{_REPO}"}],
    }


def _make_feature_detail(
    owner: str = "ts", stages: dict | None = None, id: str | None = None
):
    detail = {
        "feature_name": _FEATURE_ID,
        "title": "My Feature",
        "stage": "product_spec",
        "status": "in_design",
        "owner": owner,
        "init_pr_url": None,
        "stages": stages if stages is not None else {},
    }
    if id is not None:
        detail["id"] = id
    return detail


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


@pytest.fixture(autouse=True)
def _clean_modules():
    """Remove plugins/src modules between tests to avoid cross-test pollution."""
    keys = [k for k in sys.modules if k.startswith(("plugins", "src"))]
    for k in keys:
        del sys.modules[k]
    yield
    keys = [k for k in sys.modules if k.startswith(("plugins", "src"))]
    for k in keys:
        del sys.modules[k]


@pytest.fixture(autouse=True)
def _clear_env(monkeypatch):
    monkeypatch.delenv("GITNEXUS_MCP_URL", raising=False)
    monkeypatch.delenv("RAG_MCP_URL", raising=False)
    monkeypatch.delenv("WORKFLOW_BACKEND_URL", raising=False)
    monkeypatch.delenv("WORKFLOW_BACKEND_SERVICE_TOKEN", raising=False)
    monkeypatch.delenv("STORAGE_SERVICE_URL", raising=False)
    monkeypatch.delenv("STORAGE_SERVICE_TOKEN", raising=False)
    monkeypatch.delenv("GITHUB_TOKEN", raising=False)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _ensure_plugins_pkg():
    if "plugins" not in sys.modules:
        pkg = types.ModuleType("plugins")
        pkg.__path__ = [str(REPO_ROOT / "plugins")]
        pkg.__package__ = "plugins"
        sys.modules["plugins"] = pkg


def _sync_run_async(coro):
    """Synchronously resolve an async coroutine."""
    if asyncio.iscoroutine(coro):
        loop = asyncio.new_event_loop()
        try:
            return loop.run_until_complete(coro)
        finally:
            loop.close()
    return coro


def _load_module_file(name: str, file_path: Path) -> types.ModuleType:
    """Load a module directly from its file, bypassing package __init__."""
    _ensure_plugins_pkg()
    spec = importlib.util.spec_from_file_location(name, file_path)
    mod = importlib.util.module_from_spec(spec)
    mod.__package__ = ".".join(name.split(".")[:-1])
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod


def _patch_module_wbc(
    mod: types.ModuleType, feature_detail: dict, workspace_ctx: dict | None = None
) -> None:
    """Patch workflow_backend_client calls on a loaded module object.

    For modules that import via ``from src.services.workflow_backend_client import ...``
    at module level, patching the module attribute works.  For modules that do
    *local* imports inside function bodies (like _write_artifact in artifacts.py),
    we must patch the source module directly so the local ``from ... import``
    picks up the mock.
    """
    if workspace_ctx is None:
        workspace_ctx = _make_workspace_context()

    # Patch the source module so local-import callsites pick up the mocks.
    import sys

    wbc_name = "src.services.workflow_backend_client"
    if wbc_name not in sys.modules:
        wbc_mod = types.ModuleType(wbc_name)
        # Ensure parent packages exist
        for parent in ("src", "src.services"):
            if parent not in sys.modules:
                sys.modules[parent] = types.ModuleType(parent)
        sys.modules[wbc_name] = wbc_mod
    else:
        wbc_mod = sys.modules[wbc_name]

    wbc_mod.get_feature_detail = AsyncMock(return_value=feature_detail)
    wbc_mod.get_workspace_context = AsyncMock(return_value=workspace_ctx)
    wbc_mod.run_async = _sync_run_async

    # Also patch module-level attributes for modules that import at the top.
    mod.get_feature_detail = wbc_mod.get_feature_detail
    mod.get_workspace_context = wbc_mod.get_workspace_context
    mod.run_async = _sync_run_async


def _make_fake_ssc(*, read_return=None, write_return=None, error=None):
    """Build a fake storage_service_client module."""

    class _Error(Exception):
        def __init__(self, msg="", *, reason_code="", status=0):
            super().__init__(msg)
            self.reason_code = reason_code
            self.status = status

    fake = types.ModuleType("plugins.clients.storage_service_client")
    fake.StorageServiceError = _Error

    if error is not None:
        fake.read_document_content = MagicMock(side_effect=error)
        fake.write_document_content = MagicMock(side_effect=error)
    else:
        fake.read_document_content = MagicMock(
            return_value=read_return or {"content": "", "version_id": None}
        )
        fake.write_document_content = MagicMock(
            return_value=write_return or {"ok": True, "version_id": _VERSION_ID}
        )

    sys.modules["plugins.clients.storage_service_client"] = fake
    return fake


def _inject_ssc_into_mod(mod: types.ModuleType, fake_ssc: types.ModuleType) -> None:
    """Replace storage_service_client references on the loaded module."""
    mod.StorageServiceError = fake_ssc.StorageServiceError
    mod.read_document_content = fake_ssc.read_document_content
    if hasattr(fake_ssc, "write_document_content"):
        mod.write_document_content = fake_ssc.write_document_content


def _enter_patches(*patches_list) -> ExitStack:
    """Enter all given patch context managers and return the ExitStack."""
    stack = ExitStack()
    for p in patches_list:
        stack.enter_context(p)
    return stack


# ---------------------------------------------------------------------------
# read_file — go-owned path
# ---------------------------------------------------------------------------


class TestReadDocumentGoOwner:
    def _load(self, feature_detail: dict, ssc_read_return=None, ssc_error=None):
        fake_ssc = _make_fake_ssc(read_return=ssc_read_return, error=ssc_error)
        mod = _load_module_file(
            "plugins.tools.read", REPO_ROOT / "plugins" / "tools" / "read.py"
        )
        _patch_module_wbc(mod, feature_detail)
        _inject_ssc_into_mod(mod, fake_ssc)
        return mod, fake_ssc

    def test_go_owned_product_spec_proxies_to_storage_service(self):
        """go-owned feature: read_file proxies to storage-service, not git."""
        read_return = {"content": _PRODUCT_SPEC_CONTENT, "version_id": _VERSION_ID}
        mod, fake_ssc = self._load(
            _make_feature_detail(owner="go"), ssc_read_return=read_return
        )

        with _enter_patches(
            patch("plugins.context.get_workspace_id", return_value=_WORKSPACE_ID),
            patch("plugins.context.get_feature_id", return_value=_FEATURE_ID),
            patch("plugins.context.get_user_id", return_value="user-1"),
            patch("plugins.context.get_org_id", return_value="org-1"),
            patch("plugins.context.mark_context_gathered"),
        ):
            result = mod.handle_read_file(
                document="product_spec",
                workspace_id=_WORKSPACE_ID,
                feature_id=_FEATURE_ID,
            )

        assert result["ok"] is True
        assert result["content"] == _PRODUCT_SPEC_CONTENT
        assert result["exists"] is True
        assert result["sha"] == _VERSION_ID
        fake_ssc.read_document_content.assert_called_once_with(
            _WORKSPACE_ID,
            _FEATURE_ID,
            "product_spec.md",
            user_id="user-1",
            org_id="org-1",
        )

    def test_go_owned_technical_design_proxies_to_storage_service(self):
        """go-owned feature: read_file for technical_design proxies to storage-service."""
        read_return = {"content": _TD_CONTENT, "version_id": _VERSION_ID}
        mod, fake_ssc = self._load(
            _make_feature_detail(owner="go"), ssc_read_return=read_return
        )

        with _enter_patches(
            patch("plugins.context.get_workspace_id", return_value=_WORKSPACE_ID),
            patch("plugins.context.get_feature_id", return_value=_FEATURE_ID),
            patch("plugins.context.get_user_id", return_value="user-1"),
            patch("plugins.context.get_org_id", return_value="org-1"),
            patch("plugins.context.mark_context_gathered"),
        ):
            result = mod.handle_read_file(
                document="technical_design",
                workspace_id=_WORKSPACE_ID,
                feature_id=_FEATURE_ID,
            )

        assert result["ok"] is True
        assert result["content"] == _TD_CONTENT
        fake_ssc.read_document_content.assert_called_once_with(
            _WORKSPACE_ID,
            _FEATURE_ID,
            "tech_design.md",
            user_id="user-1",
            org_id="org-1",
        )

    def test_go_owned_not_found_returns_empty_content(self):
        """go-owned feature: 404 from storage-service → exists=False, content empty."""
        read_return = {"content": "", "version_id": None}
        mod, _fake_ssc = self._load(
            _make_feature_detail(owner="go"), ssc_read_return=read_return
        )

        with _enter_patches(
            patch("plugins.context.get_workspace_id", return_value=_WORKSPACE_ID),
            patch("plugins.context.get_feature_id", return_value=_FEATURE_ID),
            patch("plugins.context.get_user_id", return_value="user-1"),
            patch("plugins.context.get_org_id", return_value="org-1"),
            patch("plugins.context.mark_context_gathered"),
        ):
            result = mod.handle_read_file(
                document="product_spec",
                workspace_id=_WORKSPACE_ID,
                feature_id=_FEATURE_ID,
            )

        assert result["ok"] is True
        assert result["exists"] is False
        assert result["content"] == ""

    def test_status_reads_from_workflow_backend(self, monkeypatch):
        """status is synthesized from get_feature_detail's DB fields — no
        storage-service call (status.yaml doesn't exist as a document)."""
        stages = {"product_spec": {"review_status": "draft"}}
        mod, fake_ssc = self._load(_make_feature_detail(owner="go", stages=stages))

        with _enter_patches(
            patch("plugins.context.get_workspace_id", return_value=_WORKSPACE_ID),
            patch("plugins.context.get_feature_id", return_value=_FEATURE_ID),
            patch("plugins.context.get_user_id", return_value="user-1"),
            patch("plugins.context.get_org_id", return_value="org-1"),
            patch("plugins.context.mark_context_gathered"),
        ):
            result = mod.handle_read_file(
                document="status",
                workspace_id=_WORKSPACE_ID,
                feature_id=_FEATURE_ID,
            )

        assert result["ok"] is True, result.get("error")
        assert result["exists"] is True
        assert result["branch"] is None
        # storage-service must NOT have been called (status isn't stored there either)
        fake_ssc.read_document_content.assert_not_called()
        # content reflects the DB-sourced fields
        assert "in_design" in result["content"]
        assert "product_spec" in result["content"]

    def test_go_owned_storage_service_error_returns_ok_false(self):
        """go-owned feature: storage-service error → ok=False."""
        from plugins.clients.storage_service_client import StorageServiceError

        error = StorageServiceError("missing_config", reason_code="missing_config")
        mod, _fake_ssc = self._load(_make_feature_detail(owner="go"), ssc_error=error)

        with _enter_patches(
            patch("plugins.context.get_workspace_id", return_value=_WORKSPACE_ID),
            patch("plugins.context.get_feature_id", return_value=_FEATURE_ID),
            patch("plugins.context.get_user_id", return_value="user-1"),
            patch("plugins.context.get_org_id", return_value="org-1"),
            patch("plugins.context.mark_context_gathered"),
        ):
            result = mod.handle_read_file(
                document="product_spec",
                workspace_id=_WORKSPACE_ID,
                feature_id=_FEATURE_ID,
            )

        assert result["ok"] is False
        assert "error" in result

    def test_go_owned_arbitrary_filename_proxies_to_storage_service(self):
        """go-owned feature: an arbitrary filename (not one of the 3 canonical
        documents, e.g. README.md) is treated as a literal storage-service path
        instead of being rejected — regression for the old strict enum that
        made read_file unable to read anything but product_spec/
        technical_design/status."""
        read_return = {"content": "# README\n", "version_id": _VERSION_ID}
        mod, fake_ssc = self._load(
            _make_feature_detail(owner="go"), ssc_read_return=read_return
        )

        with _enter_patches(
            patch("plugins.context.get_workspace_id", return_value=_WORKSPACE_ID),
            patch("plugins.context.get_feature_id", return_value=_FEATURE_ID),
            patch("plugins.context.get_user_id", return_value="user-1"),
            patch("plugins.context.get_org_id", return_value="org-1"),
            patch("plugins.context.mark_context_gathered"),
        ):
            result = mod.handle_read_file(
                document="README.md",
                workspace_id=_WORKSPACE_ID,
                feature_id=_FEATURE_ID,
            )

        assert result["ok"] is True
        assert result["content"] == "# README\n"
        fake_ssc.read_document_content.assert_called_once_with(
            _WORKSPACE_ID,
            _FEATURE_ID,
            "README.md",
            user_id="user-1",
            org_id="org-1",
        )


# ---------------------------------------------------------------------------
# write_product_spec
# ---------------------------------------------------------------------------


class TestWriteProductSpecGoOwner:
    def _load_go(self):
        fake_ssc = _make_fake_ssc(write_return={"ok": True, "version_id": _VERSION_ID})
        mod = _load_module_file(
            "plugins.tools.artifacts", REPO_ROOT / "plugins" / "tools" / "artifacts.py"
        )
        _patch_module_wbc(mod, _make_feature_detail(owner="go"))
        _inject_ssc_into_mod(mod, fake_ssc)
        return mod, fake_ssc

    def test_proxies_to_storage_service(self):
        """write_product_spec writes to storage-service."""
        mod, fake_ssc = self._load_go()

        with _enter_patches(
            patch("plugins.context.get_workspace_id", return_value=_WORKSPACE_ID),
            patch("plugins.context.get_feature_id", return_value=_FEATURE_ID),
            patch("plugins.context.get_user_id", return_value="user-1"),
            patch("plugins.context.get_org_id", return_value="org-1"),
            patch("plugins.context.was_context_gathered", return_value=True),
        ):
            result = mod.handle_write_product_spec(
                content=_PRODUCT_SPEC_CONTENT,
                workspace_id=_WORKSPACE_ID,
                feature_id=_FEATURE_ID,
            )

        assert result["ok"] is True
        assert result.get("version_id") == _VERSION_ID
        fake_ssc.write_document_content.assert_called_once_with(
            _WORKSPACE_ID,
            _FEATURE_ID,
            "product_spec.md",
            _PRODUCT_SPEC_CONTENT,
            user_id="user-1",
            org_id="org-1",
            feature_slug=_FEATURE_ID,
        )

    def test_slug_feature_id_resolved_to_canonical_uuid(self):
        """A caller passing the feature's slug (not its UUID) as feature_id must
        not create a duplicate document — write_document_content should be
        called with the resolved UUID from get_feature_detail, not the raw slug.
        """
        fake_ssc = _make_fake_ssc(write_return={"ok": True, "version_id": _VERSION_ID})
        mod = _load_module_file(
            "plugins.tools.artifacts", REPO_ROOT / "plugins" / "tools" / "artifacts.py"
        )
        _feature_uuid = "992495d6-bfa1-4fc1-9027-16df405274ee"
        _patch_module_wbc(mod, _make_feature_detail(owner="go", id=_feature_uuid))
        _inject_ssc_into_mod(mod, fake_ssc)

        with _enter_patches(
            patch("plugins.context.get_workspace_id", return_value=_WORKSPACE_ID),
            patch("plugins.context.get_feature_id", return_value=_FEATURE_ID),
            patch("plugins.context.get_user_id", return_value="user-1"),
            patch("plugins.context.get_org_id", return_value="org-1"),
            patch("plugins.context.was_context_gathered", return_value=True),
        ):
            result = mod.handle_write_product_spec(
                content=_PRODUCT_SPEC_CONTENT,
                workspace_id=_WORKSPACE_ID,
                feature_id=_FEATURE_ID,  # the slug, not the UUID
            )

        assert result["ok"] is True
        fake_ssc.write_document_content.assert_called_once_with(
            _WORKSPACE_ID,
            _feature_uuid,
            "product_spec.md",
            _PRODUCT_SPEC_CONTENT,
            user_id="user-1",
            org_id="org-1",
            feature_slug=_FEATURE_ID,
        )

    def test_go_owned_missing_config_returns_error(self):
        """go-owned feature without storage-service config → ok=False."""
        from plugins.clients.storage_service_client import StorageServiceError

        error = StorageServiceError(
            "STORAGE_SERVICE_URL and STORAGE_SERVICE_TOKEN must both be set",
            reason_code="missing_config",
        )
        fake_ssc = _make_fake_ssc(error=error)
        mod = _load_module_file(
            "plugins.tools.artifacts", REPO_ROOT / "plugins" / "tools" / "artifacts.py"
        )
        _patch_module_wbc(mod, _make_feature_detail(owner="go"))
        _inject_ssc_into_mod(mod, fake_ssc)

        with _enter_patches(
            patch("plugins.context.get_workspace_id", return_value=_WORKSPACE_ID),
            patch("plugins.context.get_feature_id", return_value=_FEATURE_ID),
            patch("plugins.context.get_user_id", return_value="user-1"),
            patch("plugins.context.get_org_id", return_value="org-1"),
            patch("plugins.context.was_context_gathered", return_value=True),
        ):
            result = mod.handle_write_product_spec(
                content=_PRODUCT_SPEC_CONTENT,
                workspace_id=_WORKSPACE_ID,
                feature_id=_FEATURE_ID,
            )

        assert result["ok"] is False
        assert "STORAGE_SERVICE" in result["error"]


# ---------------------------------------------------------------------------
# write_technical_design — go-owned path
# ---------------------------------------------------------------------------


class TestWriteTechnicalDesignGoOwner:
    def test_proxies_to_storage_service(self):
        """write_technical_design writes to storage-service."""
        fake_ssc = _make_fake_ssc(write_return={"ok": True, "version_id": _VERSION_ID})
        mod = _load_module_file(
            "plugins.tools.artifacts", REPO_ROOT / "plugins" / "tools" / "artifacts.py"
        )
        _patch_module_wbc(mod, _make_feature_detail(owner="go"))
        _inject_ssc_into_mod(mod, fake_ssc)

        with _enter_patches(
            patch("plugins.context.get_workspace_id", return_value=_WORKSPACE_ID),
            patch("plugins.context.get_feature_id", return_value=_FEATURE_ID),
            patch("plugins.context.get_user_id", return_value="user-1"),
            patch("plugins.context.get_org_id", return_value="org-1"),
            patch("plugins.context.was_context_gathered", return_value=True),
        ):
            result = mod.handle_write_technical_design(
                content=_TD_CONTENT,
                workspace_id=_WORKSPACE_ID,
                feature_id=_FEATURE_ID,
            )

        assert result["ok"] is True
        assert result.get("version_id") == _VERSION_ID
        fake_ssc.write_document_content.assert_called_once_with(
            _WORKSPACE_ID,
            _FEATURE_ID,
            "tech_design.md",
            _TD_CONTENT,
            user_id="user-1",
            org_id="org-1",
            feature_slug=_FEATURE_ID,
        )


# ---------------------------------------------------------------------------
# edit_document
# ---------------------------------------------------------------------------


class TestEditDocumentGoOwner:
    def _load_go(self, original_content: str = "# Spec\n\nOld content.\n"):
        fake_ssc = _make_fake_ssc(
            read_return={"content": original_content, "version_id": "v0"},
            write_return={"ok": True, "version_id": _VERSION_ID},
        )
        mod = _load_module_file(
            "plugins.tools.edit", REPO_ROOT / "plugins" / "tools" / "edit.py"
        )
        _patch_module_wbc(mod, _make_feature_detail(owner="go"))
        _inject_ssc_into_mod(mod, fake_ssc)
        return mod, fake_ssc

    def test_proxies_read_apply_write_to_storage_service(self):
        """edit_document reads from, applies, and writes to storage-service."""
        original = "# Spec\n\nOld content.\n"
        edited = "# Spec\n\nNew content.\n"
        mod, fake_ssc = self._load_go(original_content=original)

        with _enter_patches(
            patch("plugins.context.get_workspace_id", return_value=_WORKSPACE_ID),
            patch("plugins.context.get_feature_id", return_value=_FEATURE_ID),
            patch("plugins.context.get_user_id", return_value="user-1"),
            patch("plugins.context.get_org_id", return_value="org-1"),
        ):
            result = mod.handle_edit_document(
                document="product_spec",
                edits=[{"old_string": "Old content.", "new_string": "New content."}],
                workspace_id=_WORKSPACE_ID,
                feature_id=_FEATURE_ID,
            )

        assert result["ok"] is True
        assert result.get("version_id") == _VERSION_ID
        fake_ssc.write_document_content.assert_called_once_with(
            _WORKSPACE_ID,
            _FEATURE_ID,
            "product_spec.md",
            edited,
            user_id="user-1",
            org_id="org-1",
        )

    def test_go_owned_edit_warns_on_missing_old_string(self):
        """go-owned edit_document: missing old_string → warning, write proceeds."""
        mod, _fake_ssc = self._load_go()

        with _enter_patches(
            patch("plugins.context.get_workspace_id", return_value=_WORKSPACE_ID),
            patch("plugins.context.get_feature_id", return_value=_FEATURE_ID),
            patch("plugins.context.get_user_id", return_value="user-1"),
            patch("plugins.context.get_org_id", return_value="org-1"),
        ):
            result = mod.handle_edit_document(
                document="product_spec",
                edits=[{"old_string": "NOT FOUND", "new_string": "replacement"}],
                workspace_id=_WORKSPACE_ID,
                feature_id=_FEATURE_ID,
            )

        assert result["ok"] is True
        assert "warnings" in result
        assert any("NOT FOUND" in w for w in result["warnings"])

    def test_go_owned_missing_config_returns_error(self):
        """go-owned edit_document without storage-service config → ok=False."""
        from plugins.clients.storage_service_client import StorageServiceError

        error = StorageServiceError(
            "STORAGE_SERVICE_URL and STORAGE_SERVICE_TOKEN must both be set",
            reason_code="missing_config",
        )
        fake_ssc = _make_fake_ssc(error=error)
        mod = _load_module_file(
            "plugins.tools.edit", REPO_ROOT / "plugins" / "tools" / "edit.py"
        )
        _patch_module_wbc(mod, _make_feature_detail(owner="go"))
        _inject_ssc_into_mod(mod, fake_ssc)

        with _enter_patches(
            patch("plugins.context.get_workspace_id", return_value=_WORKSPACE_ID),
            patch("plugins.context.get_feature_id", return_value=_FEATURE_ID),
            patch("plugins.context.get_user_id", return_value="user-1"),
            patch("plugins.context.get_org_id", return_value="org-1"),
        ):
            result = mod.handle_edit_document(
                document="product_spec",
                edits=[{"old_string": "old", "new_string": "new"}],
                workspace_id=_WORKSPACE_ID,
                feature_id=_FEATURE_ID,
            )

        assert result["ok"] is False
        assert "STORAGE_SERVICE" in result["error"]

    def test_no_stray_git_file_created(self):
        """edit_document only writes to storage-service — no git file."""
        original = "# Spec\n\nContent.\n"
        mod, fake_ssc = self._load_go(original_content=original)

        with _enter_patches(
            patch("plugins.context.get_workspace_id", return_value=_WORKSPACE_ID),
            patch("plugins.context.get_feature_id", return_value=_FEATURE_ID),
            patch("plugins.context.get_user_id", return_value="user-1"),
            patch("plugins.context.get_org_id", return_value="org-1"),
        ):
            result = mod.handle_edit_document(
                document="product_spec",
                edits=[{"old_string": "Content.", "new_string": "Updated content."}],
                workspace_id=_WORKSPACE_ID,
                feature_id=_FEATURE_ID,
            )

        assert result["ok"] is True
        fake_ssc.write_document_content.assert_called_once()


# ---------------------------------------------------------------------------
# storage_service_client unit tests (HTTP-level)
# ---------------------------------------------------------------------------


class TestStorageServiceClient:
    def test_read_document_content_happy_path(self, monkeypatch, requests_mock):
        monkeypatch.setenv("STORAGE_SERVICE_URL", _STORAGE_URL)
        monkeypatch.setenv("STORAGE_SERVICE_TOKEN", _STORAGE_TOKEN)

        requests_mock.get(
            f"{_STORAGE_URL}/api/workspaces/ws1/features/feat1/documents/content?path=product_spec.md",
            json={"content": "# Hello\n", "version_id": "v1"},
        )

        from plugins.clients.storage_service_client import read_document_content

        result = read_document_content("ws1", "feat1", "product_spec.md")
        assert result["content"] == "# Hello\n"
        assert result["version_id"] == "v1"

    def test_read_document_content_not_found(self, monkeypatch, requests_mock):
        monkeypatch.setenv("STORAGE_SERVICE_URL", _STORAGE_URL)
        monkeypatch.setenv("STORAGE_SERVICE_TOKEN", _STORAGE_TOKEN)

        requests_mock.get(
            f"{_STORAGE_URL}/api/workspaces/ws1/features/feat1/documents/content?path=product_spec.md",
            status_code=404,
        )

        from plugins.clients.storage_service_client import read_document_content

        result = read_document_content("ws1", "feat1", "product_spec.md")
        assert result["content"] == ""
        assert result["version_id"] is None

    def test_write_document_content_happy_path(self, monkeypatch, requests_mock):
        monkeypatch.setenv("STORAGE_SERVICE_URL", _STORAGE_URL)
        monkeypatch.setenv("STORAGE_SERVICE_TOKEN", _STORAGE_TOKEN)

        requests_mock.put(
            f"{_STORAGE_URL}/api/workspaces/ws1/features/feat1/documents/content?path=product_spec.md",
            json={"ok": True, "version_id": "v2"},
        )

        from plugins.clients.storage_service_client import write_document_content

        result = write_document_content(
            "ws1", "feat1", "product_spec.md", "# New content\n"
        )
        assert result["ok"] is True
        assert result["version_id"] == "v2"

    def test_missing_config_raises_error(self, monkeypatch):
        monkeypatch.delenv("STORAGE_SERVICE_URL", raising=False)
        monkeypatch.delenv("STORAGE_SERVICE_TOKEN", raising=False)

        from plugins.clients.storage_service_client import (
            StorageServiceError,
            read_document_content,
        )

        with pytest.raises(StorageServiceError) as exc_info:
            read_document_content("ws1", "feat1", "product_spec.md")
        assert exc_info.value.reason_code == "missing_config"

    def test_server_error_raises_storage_service_error(
        self, monkeypatch, requests_mock
    ):
        monkeypatch.setenv("STORAGE_SERVICE_URL", _STORAGE_URL)
        monkeypatch.setenv("STORAGE_SERVICE_TOKEN", _STORAGE_TOKEN)

        requests_mock.get(
            f"{_STORAGE_URL}/api/workspaces/ws1/features/feat1/documents/content?path=product_spec.md",
            status_code=500,
            json={"error": "internal_server_error"},
        )

        from plugins.clients.storage_service_client import (
            StorageServiceError,
            read_document_content,
        )

        with pytest.raises(StorageServiceError) as exc_info:
            read_document_content("ws1", "feat1", "product_spec.md")
        assert exc_info.value.status == 500

    def test_write_sends_content_in_body(self, monkeypatch, requests_mock):
        monkeypatch.setenv("STORAGE_SERVICE_URL", _STORAGE_URL)
        monkeypatch.setenv("STORAGE_SERVICE_TOKEN", _STORAGE_TOKEN)

        requests_mock.put(
            f"{_STORAGE_URL}/api/workspaces/ws1/features/feat1/documents/content?path=tech_design.md",
            json={"ok": True, "version_id": "v3"},
        )

        from plugins.clients.storage_service_client import write_document_content

        result = write_document_content(
            "ws1", "feat1", "tech_design.md", "# Tech Design\n"
        )
        assert result["ok"] is True

        import json

        put_body = json.loads(requests_mock.last_request.text)
        assert put_body["content"] == "# Tech Design\n"

    def test_download_image_happy_path(self, monkeypatch, requests_mock):
        monkeypatch.setenv("STORAGE_SERVICE_URL", _STORAGE_URL)
        monkeypatch.setenv("STORAGE_SERVICE_TOKEN", _STORAGE_TOKEN)

        requests_mock.get(
            f"{_STORAGE_URL}/api/workspaces/ws1/images/img-1",
            content=b"fake-png-bytes",
            headers={"Content-Type": "image/png"},
        )

        from plugins.clients.storage_service_client import download_image

        result = download_image("ws1", "img-1", user_id="user-1", org_id="org-1")
        assert result["data"] == b"fake-png-bytes"
        assert result["content_type"] == "image/png"

        req = requests_mock.last_request
        assert req.headers["Authorization"] == f"Bearer {_STORAGE_TOKEN}"
        assert req.headers["X-User-Id"] == "user-1"
        assert req.headers["X-Org-Id"] == "org-1"

    def test_download_image_not_found(self, monkeypatch, requests_mock):
        monkeypatch.setenv("STORAGE_SERVICE_URL", _STORAGE_URL)
        monkeypatch.setenv("STORAGE_SERVICE_TOKEN", _STORAGE_TOKEN)

        requests_mock.get(
            f"{_STORAGE_URL}/api/workspaces/ws1/images/missing", status_code=404
        )

        from plugins.clients.storage_service_client import (
            StorageServiceError,
            download_image,
        )

        with pytest.raises(StorageServiceError) as exc_info:
            download_image("ws1", "missing")
        assert exc_info.value.status == 404
        assert exc_info.value.reason_code == "not_found"

    def test_download_image_missing_config_raises_error(self, monkeypatch):
        monkeypatch.delenv("STORAGE_SERVICE_URL", raising=False)
        monkeypatch.delenv("STORAGE_SERVICE_TOKEN", raising=False)

        from plugins.clients.storage_service_client import (
            StorageServiceError,
            download_image,
        )

        with pytest.raises(StorageServiceError) as exc_info:
            download_image("ws1", "img-1")
        assert exc_info.value.reason_code == "missing_config"


# ---------------------------------------------------------------------------
# download_file unit tests (HTTP-level)
# ---------------------------------------------------------------------------


class TestDownloadFile:
    """Unit tests for :func:`download_file` in storage_service_client."""

    def test_download_file_success(self, monkeypatch, requests_mock):
        monkeypatch.setenv("STORAGE_SERVICE_URL", _STORAGE_URL)
        monkeypatch.setenv("STORAGE_SERVICE_TOKEN", _STORAGE_TOKEN)

        requests_mock.get(
            f"{_STORAGE_URL}/api/workspaces/ws1/files/file-1",
            content=b"fake-file-bytes",
            headers={
                "Content-Type": "application/pdf",
                "Content-Disposition": 'attachment; filename="report.pdf"',
            },
        )

        from plugins.clients.storage_service_client import download_file

        result = download_file("ws1", "file-1", user_id="user-1", org_id="org-1")
        assert result["data"] == b"fake-file-bytes"
        assert result["content_type"] == "application/pdf"
        assert result["filename"] == "report.pdf"

        req = requests_mock.last_request
        assert req.headers["Authorization"] == f"Bearer {_STORAGE_TOKEN}"
        assert req.headers["X-User-Id"] == "user-1"
        assert req.headers["X-Org-Id"] == "org-1"

    def test_download_file_not_found(self, monkeypatch, requests_mock):
        monkeypatch.setenv("STORAGE_SERVICE_URL", _STORAGE_URL)
        monkeypatch.setenv("STORAGE_SERVICE_TOKEN", _STORAGE_TOKEN)

        requests_mock.get(
            f"{_STORAGE_URL}/api/workspaces/ws1/files/missing",
            status_code=404,
        )

        from plugins.clients.storage_service_client import (
            StorageServiceError,
            download_file,
        )

        with pytest.raises(StorageServiceError) as exc_info:
            download_file("ws1", "missing")
        assert exc_info.value.status == 404
        assert exc_info.value.reason_code == "not_found"

    def test_download_file_missing_config_raises_error(self, monkeypatch):
        monkeypatch.delenv("STORAGE_SERVICE_URL", raising=False)
        monkeypatch.delenv("STORAGE_SERVICE_TOKEN", raising=False)

        from plugins.clients.storage_service_client import (
            StorageServiceError,
            download_file,
        )

        with pytest.raises(StorageServiceError) as exc_info:
            download_file("ws1", "file-1")
        assert exc_info.value.reason_code == "missing_config"


# ---------------------------------------------------------------------------
# read_uploaded_file tool tests
# ---------------------------------------------------------------------------


class TestReadUploadedFile:
    """Unit tests for the read_uploaded_file tool handler."""

    # -- helpers ----------------------------------------------------------

    @staticmethod
    def _make_pdf_bytes(text: str = "Hello PDF world!") -> bytes:
        """Create a minimal valid PDF in memory readable by PyPDF2.

        Constructs a valid PDF 1.4 document with a single page containing
        the given text in a content stream.  PyPDF2's PdfReader can parse
        this and extract_text() will return the embedded text.
        """
        # Encode the text content for the PDF stream.
        text_bytes = text.encode("ascii", errors="replace")

        # Build the content stream (a minimal PDF page description).
        stream = b"BT\n/F1 12 Tf\n72 720 Td\n(" + text_bytes + b") Tj\nET"

        # Build a minimal valid PDF by hand with correct xref offsets.
        header = b"%PDF-1.4\n"

        # Object 1: Catalog
        obj1 = b"1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n"

        # Object 2: Pages
        obj2 = b"2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n"

        # Object 3: Page
        obj3 = (
            b"3 0 obj\n"
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Contents 4 0 R /Resources << /Font << /F1 << /Type /Font "
            b"/Subtype /Type1 /BaseFont /Helvetica >> >> >> >>\n"
            b"endobj\n"
        )

        # Object 4: Content stream
        obj4 = (
            b"4 0 obj\n"
            b"<< /Length " + str(len(stream)).encode() + b" >>\n"
            b"stream\n" + stream + b"\n"
            b"endstream\n"
            b"endobj\n"
        )

        # Compute xref offsets precisely.
        offset1 = len(header)
        offset2 = offset1 + len(obj1)
        offset3 = offset2 + len(obj2)
        offset4 = offset3 + len(obj3)

        xref = (
            b"xref\n"
            b"0 5\n"
            b"0000000000 65535 f \n"
            + f"{offset1:010d}".encode()
            + b" 00000 n \n"
            + f"{offset2:010d}".encode()
            + b" 00000 n \n"
            + f"{offset3:010d}".encode()
            + b" 00000 n \n"
            + f"{offset4:010d}".encode()
            + b" 00000 n \n"
        )

        startxref = offset1 + len(obj1) + len(obj2) + len(obj3) + len(obj4)
        trailer = (
            b"trailer\n"
            b"<< /Size 5 /Root 1 0 R >>\n"
            b"startxref\n" + str(startxref).encode() + b"\n%%EOF"
        )

        return header + obj1 + obj2 + obj3 + obj4 + xref + trailer

    @staticmethod
    def _make_docx_bytes(text: str = "Hello DOCX world!") -> bytes:
        """Create a minimal valid DOCX in memory."""
        import io

        from docx import Document as DocxDocument

        doc = DocxDocument()
        doc.add_paragraph(text)
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    @staticmethod
    def _make_xlsx_bytes(data=None) -> bytes:
        """Create a minimal valid XLSX in memory."""
        import io

        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        if data is None:
            data = [["Name", "Value"], ["Alice", "42"], ["Bob", "7"]]
        for row in data:
            ws.append(row)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    # -- download_file tests (HTTP-level) ---------------------------------

    def test_download_file_success(self, monkeypatch, requests_mock):
        monkeypatch.setenv("STORAGE_SERVICE_URL", _STORAGE_URL)
        monkeypatch.setenv("STORAGE_SERVICE_TOKEN", _STORAGE_TOKEN)

        requests_mock.get(
            f"{_STORAGE_URL}/api/workspaces/ws1/files/file-1",
            content=b"hello world",
            headers={"Content-Type": "text/plain"},
        )

        from plugins.tools.read_uploaded_file import handle

        result = handle(file_id="file-1", workspace_id="ws1")
        assert result["ok"] is True
        assert result["text"] == "hello world"
        assert result["content_type"] == "text/plain"
        assert result["truncated"] is False

    def test_download_file_not_found(self, monkeypatch, requests_mock):
        monkeypatch.setenv("STORAGE_SERVICE_URL", _STORAGE_URL)
        monkeypatch.setenv("STORAGE_SERVICE_TOKEN", _STORAGE_TOKEN)

        requests_mock.get(
            f"{_STORAGE_URL}/api/workspaces/ws1/files/missing",
            status_code=404,
        )

        from plugins.tools.read_uploaded_file import handle

        result = handle(file_id="missing", workspace_id="ws1")
        assert result["ok"] is False
        assert "not_found" in result.get("error", "").lower() or "404" in result.get(
            "error", ""
        )

    # -- format-specific parsing tests ------------------------------------

    def test_read_uploaded_file_txt(self, monkeypatch, requests_mock):
        monkeypatch.setenv("STORAGE_SERVICE_URL", _STORAGE_URL)
        monkeypatch.setenv("STORAGE_SERVICE_TOKEN", _STORAGE_TOKEN)

        requests_mock.get(
            f"{_STORAGE_URL}/api/workspaces/ws1/files/txt-1",
            content=b"Hello, this is a text file.\nLine two.",
            headers={"Content-Type": "text/plain"},
        )

        from plugins.tools.read_uploaded_file import handle

        result = handle(file_id="txt-1", workspace_id="ws1")
        assert result["ok"] is True
        assert "Hello, this is a text file." in result["text"]
        assert "Line two." in result["text"]

    def test_read_uploaded_file_pdf(self, monkeypatch, requests_mock):
        monkeypatch.setenv("STORAGE_SERVICE_URL", _STORAGE_URL)
        monkeypatch.setenv("STORAGE_SERVICE_TOKEN", _STORAGE_TOKEN)

        pdf_bytes = self._make_pdf_bytes("Hello PDF world!")
        requests_mock.get(
            f"{_STORAGE_URL}/api/workspaces/ws1/files/pdf-1",
            content=pdf_bytes,
            headers={"Content-Type": "application/pdf"},
        )

        from plugins.tools.read_uploaded_file import handle

        result = handle(file_id="pdf-1", workspace_id="ws1")
        assert result["ok"] is True
        assert "Hello PDF world!" in result["text"]

    def test_read_uploaded_file_docx(self, monkeypatch, requests_mock):
        monkeypatch.setenv("STORAGE_SERVICE_URL", _STORAGE_URL)
        monkeypatch.setenv("STORAGE_SERVICE_TOKEN", _STORAGE_TOKEN)

        docx_bytes = self._make_docx_bytes("Hello DOCX world!")
        requests_mock.get(
            f"{_STORAGE_URL}/api/workspaces/ws1/files/docx-1",
            content=docx_bytes,
            headers={
                "Content-Type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            },
        )

        from plugins.tools.read_uploaded_file import handle

        result = handle(file_id="docx-1", workspace_id="ws1")
        assert result["ok"] is True
        assert "Hello DOCX world!" in result["text"]

    def test_read_uploaded_file_xlsx(self, monkeypatch, requests_mock):
        monkeypatch.setenv("STORAGE_SERVICE_URL", _STORAGE_URL)
        monkeypatch.setenv("STORAGE_SERVICE_TOKEN", _STORAGE_TOKEN)

        xlsx_bytes = self._make_xlsx_bytes()
        requests_mock.get(
            f"{_STORAGE_URL}/api/workspaces/ws1/files/xlsx-1",
            content=xlsx_bytes,
            headers={
                "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            },
        )

        from plugins.tools.read_uploaded_file import handle

        result = handle(file_id="xlsx-1", workspace_id="ws1")
        assert result["ok"] is True
        assert "Sheet: Sheet1" in result["text"]
        assert "Alice" in result["text"]
        assert "Bob" in result["text"]

    # -- extension-based fallback tests -----------------------------------

    def test_read_uploaded_file_pdf_by_extension(self, monkeypatch, requests_mock):
        """Content-Type is generic, but filename ends in .pdf → PDF parser used."""
        monkeypatch.setenv("STORAGE_SERVICE_URL", _STORAGE_URL)
        monkeypatch.setenv("STORAGE_SERVICE_TOKEN", _STORAGE_TOKEN)

        pdf_bytes = self._make_pdf_bytes("PDF via extension")
        requests_mock.get(
            f"{_STORAGE_URL}/api/workspaces/ws1/files/pdf-ext",
            content=pdf_bytes,
            headers={
                "Content-Type": "application/octet-stream",
                "Content-Disposition": 'attachment; filename="myfile.pdf"',
            },
        )

        from plugins.tools.read_uploaded_file import handle

        result = handle(file_id="pdf-ext", workspace_id="ws1")
        assert result["ok"] is True
        assert "PDF via extension" in result["text"]

    def test_read_uploaded_file_xlsx_by_extension(self, monkeypatch, requests_mock):
        """Content-Type is generic, but filename ends in .xlsx → XLSX parser used."""
        monkeypatch.setenv("STORAGE_SERVICE_URL", _STORAGE_URL)
        monkeypatch.setenv("STORAGE_SERVICE_TOKEN", _STORAGE_TOKEN)

        xlsx_bytes = self._make_xlsx_bytes([["Col1", "Col2"], ["a", "b"]])
        requests_mock.get(
            f"{_STORAGE_URL}/api/workspaces/ws1/files/xlsx-ext",
            content=xlsx_bytes,
            headers={
                "Content-Type": "application/octet-stream",
                "Content-Disposition": 'attachment; filename="data.xlsx"',
            },
        )

        from plugins.tools.read_uploaded_file import handle

        result = handle(file_id="xlsx-ext", workspace_id="ws1")
        assert result["ok"] is True
        assert "Sheet: Sheet1" in result["text"]

    # -- error handling tests ---------------------------------------------

    def test_read_uploaded_file_corrupt(self, monkeypatch, requests_mock):
        """Random binary that is not valid PDF/DOCX/XLSX and not UTF-8 → error."""
        monkeypatch.setenv("STORAGE_SERVICE_URL", _STORAGE_URL)
        monkeypatch.setenv("STORAGE_SERVICE_TOKEN", _STORAGE_TOKEN)

        # Non-UTF-8 binary bytes
        corrupt_bytes = b"\x80\x81\x82\xff\xfe\xfd"
        requests_mock.get(
            f"{_STORAGE_URL}/api/workspaces/ws1/files/corrupt",
            content=corrupt_bytes,
            headers={"Content-Type": "application/octet-stream"},
        )

        from plugins.tools.read_uploaded_file import handle

        result = handle(file_id="corrupt", workspace_id="ws1")
        assert result["ok"] is False
        assert "error" in result

    # -- truncation tests -------------------------------------------------

    def test_read_uploaded_file_truncation(self, monkeypatch, requests_mock):
        """Content exceeding 100 KB is truncated with a marker."""
        monkeypatch.setenv("STORAGE_SERVICE_URL", _STORAGE_URL)
        monkeypatch.setenv("STORAGE_SERVICE_TOKEN", _STORAGE_TOKEN)

        # Generate text that is ~200 KB (well above 100 KB limit).
        chunk = "Hello world! This is a line of text.\n"
        big_text = chunk * 6000  # ~200 KB
        big_bytes = big_text.encode("utf-8")

        requests_mock.get(
            f"{_STORAGE_URL}/api/workspaces/ws1/files/big-file",
            content=big_bytes,
            headers={"Content-Type": "text/plain"},
        )

        from plugins.tools.read_uploaded_file import handle

        result = handle(file_id="big-file", workspace_id="ws1")
        assert result["ok"] is True
        assert result["truncated"] is True
        assert "[... truncated ...]" in result["text"]
        # Truncated text should be <= ~100 KB + marker
        assert len(result["text"].encode("utf-8")) <= 110_000

    # -- missing file_id / workspace_id -----------------------------------

    def test_read_uploaded_file_missing_file_id(self):
        from plugins.tools.read_uploaded_file import handle

        result = handle(file_id="", workspace_id="ws1")
        assert result["ok"] is False
        assert "file_id" in result.get("error", "").lower()

    def test_read_uploaded_file_missing_workspace_id(self):
        from plugins.tools.read_uploaded_file import handle

        result = handle(file_id="file-1", workspace_id="")
        assert result["ok"] is False
        assert "workspace_id" in result.get("error", "").lower()
